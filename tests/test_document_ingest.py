"""
Tests for retrieval/document_ingest.py's extract_text() -- offline, no
network, no live key. Builds real .pdf/.docx bytes in-memory with the same
pypdf/python-docx libraries the module itself uses to parse them, so this
proves round-trip correctness against the real libraries, not a mock of
them.
"""

import io

import pytest
from pypdf import PdfWriter

from anticassiopeia.retrieval.document_ingest import DocumentIngestError, SUPPORTED_EXTENSIONS, extract_text

# --- plain text formats ---


@pytest.mark.parametrize("ext", [".txt", ".md", ".csv", ".rpt", ".svd", ".xml"])
def test_plain_text_extensions_are_read_as_utf8_verbatim(ext):
    text = extract_text(f"hello world, café".encode("utf-8"), f"doc{ext}")
    assert text == "hello world, café"


def test_plain_text_with_invalid_utf8_is_a_named_error_not_a_crash():
    with pytest.raises(DocumentIngestError, match="valid UTF-8"):
        extract_text(b"\xff\xfe\x00\x01not utf8", "broken.txt")


def test_empty_plain_text_file_is_a_named_error():
    with pytest.raises(DocumentIngestError, match="empty"):
        extract_text(b"", "empty.txt")


def test_whitespace_only_plain_text_file_is_a_named_error():
    with pytest.raises(DocumentIngestError, match="empty"):
        extract_text(b"   \n\n  ", "blank.txt")


# --- unsupported extension ---


def test_unsupported_extension_is_a_named_error_listing_what_is_supported():
    with pytest.raises(DocumentIngestError, match="Unsupported file extension"):
        extract_text(b"whatever", "archive.zip")
    # the error message actually lists the real supported set, not a stale one
    try:
        extract_text(b"whatever", "archive.zip")
    except DocumentIngestError as exc:
        for ext in SUPPORTED_EXTENSIONS:
            assert ext in str(exc)


def test_no_extension_at_all_is_a_named_error_not_a_crash():
    with pytest.raises(DocumentIngestError, match="Unsupported file extension ''"):
        extract_text(b"whatever", "README")


# --- PDF, via a real pypdf-built PDF ---


def _build_real_pdf_with_text(lines: list[str]) -> bytes:
    """Builds a real, minimal, valid PDF using pypdf's own writer with an
    embedded text-drawing content stream -- round-trips through the exact
    library extract_text() itself uses, so a pass here means the real
    pypdf API contract is satisfied, not a mocked stand-in for it."""
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    content = "BT /F1 12 Tf 10 180 Td\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj 0 -14 TD\n"
    content += "ET"

    # Write the content stream directly via pypdf's low-level generic API --
    # the simplest way to get REAL, library-extractable text into a REAL PDF
    # without a heavier dependency (reportlab etc.) just for a test fixture.
    content_obj = DecodedStreamObject()
    content_obj.set_data(content.encode("latin-1"))
    content_ref = writer._add_object(content_obj)
    page[NameObject("/Contents")] = content_ref

    font_dict = DictionaryObject()
    font_dict[NameObject("/Type")] = NameObject("/Font")
    font_dict[NameObject("/Subtype")] = NameObject("/Type1")
    font_dict[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font_dict)

    resources = DictionaryObject()
    font_res = DictionaryObject()
    font_res[NameObject("/F1")] = font_ref
    resources[NameObject("/Font")] = font_res
    page[NameObject("/Resources")] = resources

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_pdf_with_a_real_text_layer_extracts_the_real_text():
    pdf_bytes = _build_real_pdf_with_text(["Maximum Rated Depth: 300 m"])
    text = extract_text(pdf_bytes, "datasheet.pdf")
    assert "Maximum Rated Depth: 300 m" in text


def test_pdf_with_no_text_layer_is_a_named_error_not_a_silent_empty_string():
    """A blank page (no content stream at all) is pypdf's own stand-in for
    a scanned/image-only PDF with no embedded text layer -- OCR is
    explicitly out of scope (see module docstring)."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(DocumentIngestError, match="no extractable text"):
        extract_text(buf.getvalue(), "scanned.pdf")


def test_malformed_pdf_bytes_are_a_named_error_not_a_crash():
    with pytest.raises(DocumentIngestError, match="could not be read as a PDF"):
        extract_text(b"%PDF-1.4 this is not actually a real pdf structure", "broken.pdf")


# --- DOCX, via a real python-docx-built document ---


def test_docx_with_real_paragraphs_extracts_the_real_text():
    import docx

    document = docx.Document()
    document.add_paragraph("The Widget Mk4 draws a maximum of 4.2 watts under full load.")
    buf = io.BytesIO()
    document.save(buf)
    text = extract_text(buf.getvalue(), "manual.docx")
    assert "The Widget Mk4 draws a maximum of 4.2 watts under full load." in text


def test_docx_with_no_paragraph_text_is_a_named_error():
    import docx

    document = docx.Document()  # no paragraphs added at all
    buf = io.BytesIO()
    document.save(buf)
    with pytest.raises(DocumentIngestError, match="no extractable text"):
        extract_text(buf.getvalue(), "empty.docx")


def test_malformed_docx_bytes_are_a_named_error_not_a_crash():
    with pytest.raises(DocumentIngestError, match="could not be read as a DOCX"):
        extract_text(b"this is not a real docx/zip file at all", "broken.docx")
